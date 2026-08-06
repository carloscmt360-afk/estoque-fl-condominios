#!/usr/bin/env python3
"""Extrai o histórico de custo por departamento da planilha do usuário para o
formato que `core-cli --import-history` consome.

Duas fontes, porque a planilha tem duas:

* **2026 (e qualquer ano com detalhe)** sai da aba ``CUSTO POR DEPTO``, que é
  a origem real: a aba ``RESTROSPECTO`` só a resume com SUMIFS.
* **Anos sem detalhe (2025)** saem do próprio quadro da ``RESTROSPECTO``,
  onde os valores estão digitados à mão — não existe outra fonte.

Por que não copiar o quadro do RESTROSPECTO também para 2026: os SUMIFS dele
perdem lançamentos em silêncio. Confirmado no arquivo de 30/07:

* o quadro lista 13 departamentos fixos, então CONTABILIDADE e DIRETORIA
  COMERCIAL não entram em lugar nenhum;
* uma linha de DP de maio está digitada como ``'DP '`` (com espaço no fim),
  e o SUMIFS não casa com ``DP``.

Este script agrega pelo mesmo par (MÊS REF., ANO REF.) que o SUMIFS usa, mas
normalizando o nome do departamento — então recupera esses lançamentos. O
total sai maior que o da planilha, e a diferença é impressa no fim para
conferência.

Uso:
    python3 extrair_retrospecto.py <planilha.xlsx> <saida.json>
"""
import json
import sys
import unicodedata
from collections import defaultdict

import openpyxl

MESES = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO",
         "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]
MES_INDICE = {m: i for i, m in enumerate(MESES)}

ABA_DETALHE = "CUSTO POR DEPTO"
ABA_QUADRO = "RESTROSPECTO"  # sic — o nome da aba tem esse typo no arquivo original


def normaliza_mes(valor):
    """'Fevereiro', 'FEVEREIRO' e 'fevereiro ' são o mesmo mês (o SUMIFS da
    planilha é insensível a caixa, e o arquivo real usa as duas grafias)."""
    chave = unicodedata.normalize("NFC", str(valor or "")).strip().upper()
    return MES_INDICE.get(chave)


def normaliza_depto(valor):
    """Mesma regra do canonDeptKey() em C++: apara, colapsa espaços internos e
    sobe para maiúsculas. O nome exibido é o normalizado, para as duas fontes
    combinarem numa linha só."""
    s = unicodedata.normalize("NFC", str(valor or "")).strip()
    return " ".join(s.split()).upper()


def do_detalhe(wb):
    """Agrega a aba de detalhe por (ano, mês, departamento)."""
    ws = wb[ABA_DETALHE]
    agregado = defaultdict(float)
    ignoradas = []
    for linha in ws.iter_rows(min_row=3, values_only=True):
        mes, ano, depto, _material, _qtd, _data, _valor, total = linha[:8]
        if depto is None or total is None:
            continue
        m = normaliza_mes(mes)
        if m is None or ano is None:
            ignoradas.append((mes, ano, depto, total))
            continue
        agregado[(int(ano), m, normaliza_depto(depto))] += float(total)
    return agregado, ignoradas


def do_quadro(wb, pular_anos):
    """Lê os quadros anuais da aba RESTROSPECTO (blocos 'ANO | JANEIRO..DEZEMBRO').

    Só é usada para anos que NÃO têm detalhe — ver docstring do módulo."""
    ws = wb[ABA_QUADRO]
    agregado = defaultdict(float)
    anos_lidos = set()
    for linha in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        # o rótulo do bloco fica na coluna B (índice 1) e é o ano
        rotulo = linha[1] if len(linha) > 1 else None
        try:
            ano = int(str(rotulo).strip())
        except (TypeError, ValueError):
            continue
        if not (2000 <= ano <= 2100):
            continue
        # cabeçalho encontrado: as linhas seguintes são os departamentos
        cabecalho_row = None
        for idx, row in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
            if len(row) > 1 and str(row[1]).strip() == str(ano) and str(row[2]).strip().upper() == "JANEIRO":
                cabecalho_row = idx
                break
        if cabecalho_row is None:
            continue
        if ano in pular_anos:
            continue
        for row in ws.iter_rows(min_row=cabecalho_row + 1, values_only=True):
            nome = row[1] if len(row) > 1 else None
            if nome is None or not str(nome).strip():
                break  # linha em branco fecha o bloco
            depto = normaliza_depto(nome)
            for m in range(12):
                v = row[2 + m] if len(row) > 2 + m else None
                if isinstance(v, (int, float)):
                    agregado[(ano, m, depto)] += float(v)
            anos_lidos.add(ano)
    return agregado, anos_lidos


def total_do_quadro(wb, ano):
    """Total anual como a planilha o exibe — a referência de conferência."""
    ws = wb[ABA_QUADRO]
    for idx, row in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
        if len(row) > 2 and str(row[1]).strip() == str(ano) and str(row[2]).strip().upper() == "JANEIRO":
            total = 0.0
            for r in ws.iter_rows(min_row=idx + 1, values_only=True):
                nome = r[1] if len(r) > 1 else None
                if nome is None or not str(nome).strip():
                    break
                for m in range(12):
                    v = r[2 + m] if len(r) > 2 + m else None
                    if isinstance(v, (int, float)):
                        total += float(v)
            return total
    return None


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    origem, destino = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(origem, data_only=True)

    detalhe, ignoradas = do_detalhe(wb)
    anos_com_detalhe = {ano for (ano, _, _) in detalhe}
    quadro, anos_do_quadro = do_quadro(wb, pular_anos=anos_com_detalhe)

    agregado = defaultdict(float)
    for chave, v in detalhe.items():
        agregado[chave] += v
    for chave, v in quadro.items():
        agregado[chave] += v

    linhas = [
        {"year": ano, "month0": mes, "dept": depto, "amount": round(v, 2)}
        for (ano, mes, depto), v in sorted(agregado.items())
        if abs(v) >= 0.005
    ]
    with open(destino, "w", encoding="utf-8") as fh:
        json.dump({"rows": linhas}, fh, ensure_ascii=False, indent=1)

    print(f"gravado: {destino} ({len(linhas)} linhas)")
    print(f"  anos do detalhe ({ABA_DETALHE}): {sorted(anos_com_detalhe)}")
    print(f"  anos do quadro  ({ABA_QUADRO}) : {sorted(anos_do_quadro)}")
    if ignoradas:
        print(f"  ATENÇÃO: {len(ignoradas)} linha(s) do detalhe sem MÊS/ANO REF. válidos foram ignoradas")

    por_ano = defaultdict(float)
    for l in linhas:
        por_ano[l["year"]] += l["amount"]
    print("\n  ano   importado      planilha   diferença")
    for ano in sorted(por_ano):
        planilha = total_do_quadro(wb, ano)
        if planilha is None:
            print(f"  {ano} {por_ano[ano]:>12.2f}  (sem quadro na planilha)")
        else:
            print(f"  {ano} {por_ano[ano]:>12.2f} {planilha:>13.2f} {por_ano[ano] - planilha:>11.2f}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
